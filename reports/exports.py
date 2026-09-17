import csv
from io import BytesIO
from decimal import Decimal
from html import escape
from django.http import HttpResponse,Http404
def safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')):return "'"+value
    return value
def export_table(title,columns,rows,fmt):
    if fmt=='csv':
        r=HttpResponse(content_type='text/csv; charset=utf-8');r.write('\ufeff');writer=csv.writer(r,delimiter=';');writer.writerow(columns)
        for row in rows:writer.writerow([safe_cell(v) for v in row])
    elif fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='Relatório';ws.append(columns)
        for row in rows:ws.append([float(v) if isinstance(v,Decimal) else safe_cell(v) for v in row])
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='078A82')
        for i,label in enumerate(columns,1):ws.column_dimensions[get_column_letter(i)].width=min(45,max(18,len(label)+5))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='#,##0.00'
        stream=BytesIO();wb.save(stream);r=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    elif fmt=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        stream=BytesIO();styles=getSampleStyleSheet();styles['BodyText'].fontSize=8;styles['BodyText'].leading=11
        data=[[Paragraph(escape(str(v)),styles['BodyText']) for v in row] for row in [columns,*rows]]
        t=Table(data,repeatRows=1,colWidths=[(landscape(A4)[0]-50)/len(columns)]*len(columns))
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DFF3F1')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FA')]),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        def footer(c,doc):
            c.setFont('Helvetica',8);c.drawString(25,16,'OdontoClínica · Relatório gerencial');c.drawRightString(815,16,f'Página {doc.page}')
        SimpleDocTemplate(stream,pagesize=landscape(A4),leftMargin=25,rightMargin=25,topMargin=30,bottomMargin=30).build([Paragraph(escape(title),styles['Title']),Spacer(1,12),t],onFirstPage=footer,onLaterPages=footer)
        r=HttpResponse(stream.getvalue(),content_type='application/pdf')
    else:raise Http404
    r['Content-Disposition']=f'attachment; filename="relatorio-odontoclinica.{fmt}"';r['Cache-Control']='no-store';return r
